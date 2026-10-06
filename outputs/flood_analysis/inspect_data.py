import openpyxl, json, collections, datetime
from pathlib import Path
p=Path(r'C:/Users/admin/OneDrive/กลุ่มพยากรณ์สุขภาพ/2569_น้ำท่วม/SAT/ปภ/รายงานสถานการณ์สาธารณภัยรายจังหวัด.xlsx')
w=openpyxl.load_workbook(p,read_only=True,data_only=True)
data={}
for s in w:
 rows=[(i,list(r)) for i,r in enumerate(s.iter_rows(values_only=True),1) if any(v is not None for v in r)]
 h=rows[0][1]; records=[dict(zip(h,r)) for i,r in rows[1:]]; data[s.title]=records
 print('\nSHEET',s.title,'nonempty',len(rows),'lastrow',rows[-1][0]);print('HEADERS',[(i+1,x) for i,x in enumerate(h) if x])
 if s.title=='Dashboard':
  print('ROWS',rows[:18]);continue
 for k in h:
  if not k or 'raw_' in k or 'url' in k.lower() or k in ['Remarks','Record_ID','station_name','name','District_Names']:continue
  vals=[r.get(k) for r in records]; c=collections.Counter(str(x) for x in vals)
  if len(c)<=18:print(k,dict(c))
  elif any(t in k.lower() for t in ['date','time','observed','fetched','updated','generated']): print(k,'range',min(str(x) for x in vals if x is not None),max(str(x) for x in vals if x is not None),'missing',vals.count(None))
  else: print(k,'distinct',len(c),'missing',vals.count(None),'sample',list(c)[:3])
 if s.title=='Disaster_DB':
  groups=collections.defaultdict(list)
  for r in records: groups[(str(r['Report_Date']),r['Disaster_Type'],r['Current_Status'])].append(r)
  for k,rs in groups.items():print('GROUP',k,'rows',len(rs),'provinces',len(set(r['Province'] for r in rs)),'households',sum(r['Affected_Households'] for r in rs if isinstance(r['Affected_Households'],(int,float))),'districts',sum(r['Affected_Districts_Count'] for r in rs if isinstance(r['Affected_Districts_Count'],(int,float))))
  print('DUP_IDS',[(k,v) for k,v in collections.Counter(r['Record_ID'] for r in records).items() if v>1])
 if s.title=='shelter_DB':
  print('TOTALS',{k:sum(r[k] for r in records if isinstance(r.get(k),(int,float))) for k in ['capacity','occupied','available']})
wf=openpyxl.load_workbook(p,read_only=True,data_only=False)
print('DASH_FORMULAS',[(c.coordinate,c.value) for row in wf['Dashboard'] for c in row if c.data_type=='f'][:30])
out=p.parent.parent/'outputs'/'flood_analysis'/'source_data.json'
out.write_text(json.dumps(data,ensure_ascii=False,default=lambda x:x.isoformat()),encoding='utf-8')
rs=data['thai_water_DB']; print('FUTURE',sum(r['age_minutes_at_fetch']<0 for r in rs))
ss=data['shelter_DB']; print('SHELTER_MISMATCH',[(r['district'],r['shelter_name'],r['capacity'],r['occupied'],r['available']) for r in ss if r['capacity']-r['occupied']!=r['available']])
print('BKK_ROWS', [{k:v for k,v in r.items() if k not in ['Remarks','Source_URL']} for r in data['Disaster_DB'] if r['Province']=='กรุงเทพมหานคร'])
print('BKK_AREAS',sorted(set(r['district_or_area'] for r in data['BKK_water_DB'])))
